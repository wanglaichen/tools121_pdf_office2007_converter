from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


@dataclass(frozen=True)
class ConversionResult:
    page_count: int
    dpi: int
    output_path: Path


class ConversionError(RuntimeError):
    """Raised when a PDF cannot be converted to an Office 2007 compatible DOCX."""


def convert_pdf_to_office2007_docx(
    pdf_path: str | Path,
    docx_path: str | Path,
    *,
    dpi: int = 220,
) -> ConversionResult:
    """
    Convert a PDF to DOCX by rendering each PDF page as a PNG and placing it
    into a Word page. This matches the Office 2007-compatible conversion style
    that opened correctly in Word 2007.
    """
    pdf_path = Path(pdf_path)
    docx_path = Path(docx_path)

    if dpi < 100 or dpi > 300:
        raise ConversionError("DPI must be between 100 and 300.")
    if not pdf_path.exists():
        raise ConversionError(f"PDF file not found: {pdf_path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise ConversionError("Only PDF files are supported.")

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    scale = dpi / 72

    try:
        pdf = fitz.open(str(pdf_path))
    except Exception as exc:
        raise ConversionError(f"Unable to open PDF: {exc}") from exc

    try:
        if pdf.page_count == 0:
            raise ConversionError("PDF has no pages.")

        first_rect = pdf[0].rect
        doc = Document()
        section = doc.sections[0]
        _configure_page(section, first_rect)

        page_width_inches = first_rect.width / 72
        page_height_inches = first_rect.height / 72

        for page_index, page in enumerate(pdf):
            if page_index == 0 and doc.paragraphs:
                paragraph = doc.paragraphs[0]
            else:
                if page_index > 0:
                    doc.add_page_break()
                paragraph = doc.add_paragraph()

            _configure_paragraph(paragraph)

            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image_stream = BytesIO(pixmap.tobytes("png"))

            render_width_inches = page.rect.width / 72
            render_height_inches = page.rect.height / 72
            fit_scale = min(
                page_width_inches / render_width_inches,
                page_height_inches / render_height_inches,
            )

            paragraph.add_run().add_picture(
                image_stream,
                width=Inches(render_width_inches * fit_scale),
                height=Inches(render_height_inches * fit_scale),
            )

        doc.save(str(docx_path))
        return ConversionResult(
            page_count=pdf.page_count,
            dpi=dpi,
            output_path=docx_path,
        )
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Conversion failed: {exc}") from exc
    finally:
        pdf.close()


def convert_pdf_to_editable_docx(
    pdf_path: str | Path,
    docx_path: str | Path,
) -> ConversionResult:
    """
    Convert a text-based PDF to an editable DOCX. This mode prioritizes editable
    text and approximate style preservation over pixel-perfect layout.
    """
    pdf_path = Path(pdf_path)
    docx_path = Path(docx_path)

    if not pdf_path.exists():
        raise ConversionError(f"PDF file not found: {pdf_path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise ConversionError("Only PDF files are supported.")

    docx_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        pdf = fitz.open(str(pdf_path))
    except Exception as exc:
        raise ConversionError(f"Unable to open PDF: {exc}") from exc

    try:
        if pdf.page_count == 0:
            raise ConversionError("PDF has no pages.")

        doc = Document()
        first_rect = pdf[0].rect
        _configure_editable_page(doc.sections[0], first_rect)

        extracted_chars = 0
        for page_index, page in enumerate(pdf):
            if page_index > 0:
                doc.add_page_break()

            lines = _extract_text_lines(page)
            if not lines:
                continue

            previous_bottom = 0
            for line in lines:
                paragraph = doc.add_paragraph()
                _configure_editable_paragraph(paragraph, line, previous_bottom)
                _add_line_runs(paragraph, line)
                previous_bottom = line["bbox"][3]
                extracted_chars += sum(len(span["text"]) for span in line["spans"])

        if extracted_chars == 0:
            raise ConversionError(
                "PDF does not contain extractable text. Use image mode or OCR first."
            )

        doc.save(str(docx_path))
        return ConversionResult(
            page_count=pdf.page_count,
            dpi=0,
            output_path=docx_path,
        )
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Conversion failed: {exc}") from exc
    finally:
        pdf.close()


def _configure_page(section, page_rect) -> None:
    section.page_width = Inches(page_rect.width / 72)
    section.page_height = Inches(page_rect.height / 72)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)


def _configure_paragraph(paragraph) -> None:
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    paragraph.paragraph_format.line_spacing = 1


def _configure_editable_page(section, page_rect) -> None:
    section.page_width = Inches(page_rect.width / 72)
    section.page_height = Inches(page_rect.height / 72)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)


def _extract_text_lines(page) -> list[dict]:
    raw = page.get_text("dict")
    lines = []

    for block in raw.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = []
            for span in line.get("spans", []):
                text = span.get("text", "")
                if not text:
                    continue
                spans.append(span)

            if not spans:
                continue

            bbox = line.get("bbox")
            if not bbox:
                x0 = min(span["bbox"][0] for span in spans)
                y0 = min(span["bbox"][1] for span in spans)
                x1 = max(span["bbox"][2] for span in spans)
                y1 = max(span["bbox"][3] for span in spans)
                bbox = (x0, y0, x1, y1)

            lines.append({"bbox": bbox, "spans": spans})

    lines.sort(key=lambda item: (round(item["bbox"][1], 1), item["bbox"][0]))
    return lines


def _configure_editable_paragraph(paragraph, line: dict, previous_bottom: float) -> None:
    x0, y0, _x1, y1 = line["bbox"]
    paragraph.paragraph_format.left_indent = Pt(max(0, x0))
    paragraph.paragraph_format.space_before = Pt(max(0, y0 - previous_bottom))
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(max(8, y1 - y0 + 1))


def _add_line_runs(paragraph, line: dict) -> None:
    previous_x = line["bbox"][0]
    for span in line["spans"]:
        text = span.get("text", "")
        if not text:
            continue

        gap = span["bbox"][0] - previous_x
        if gap > max(8, span.get("size", 10) * 0.6):
            paragraph.add_run(" " * max(1, round(gap / max(span.get("size", 10) * 0.35, 3))))

        run = paragraph.add_run(text)
        _apply_span_style(run, span)
        previous_x = span["bbox"][2]


def _apply_span_style(run, span: dict) -> None:
    font = run.font
    font.size = Pt(max(1, float(span.get("size", 10))))

    font_name = _font_name(span.get("font", ""))
    if font_name:
        font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    flags = int(span.get("flags", 0) or 0)
    source_font = span.get("font", "").lower()
    font.bold = bool(flags & 16) or "bold" in source_font or "black" in source_font
    font.italic = bool(flags & 2) or "italic" in source_font or "oblique" in source_font

    color = span.get("color")
    if isinstance(color, int):
        r = (color >> 16) & 255
        g = (color >> 8) & 255
        b = color & 255
        font.color.rgb = RGBColor(r, g, b)


def _font_name(pdf_font_name: str) -> str:
    normalized = pdf_font_name.replace("-", "").replace("_", "").lower()
    if "yahei" in normalized or "microsoft" in normalized:
        return "Microsoft YaHei"
    if "simsun" in normalized or "song" in normalized:
        return "SimSun"
    if "simhei" in normalized or "hei" in normalized:
        return "SimHei"
    if "kaiti" in normalized or "kai" in normalized:
        return "KaiTi"
    if "fangsong" in normalized:
        return "FangSong"
    if "arial" in normalized:
        return "Arial"
    if "times" in normalized:
        return "Times New Roman"
    if "calibri" in normalized:
        return "Calibri"
    return "Microsoft YaHei"
